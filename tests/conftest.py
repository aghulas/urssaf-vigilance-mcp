from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def pdf_exemple() -> Path:
    return FIXTURES / "attestation_exemple.pdf"


@pytest.fixture
def classeur_suivi(tmp_path: Path) -> Path:
    """Classeur minimal reproduisant la structure du tableau réel (titre, sous-titre, ligne vide, en-têtes)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Suivi attestations"
    ws["A1"] = "Suivi des attestations - Chantier test"
    ws["A2"] = "situation au 13/10/2026"
    entetes = [
        "Opération", "Lot", "Entreprise", "URSSAF\ndate attestation", "URSSAF\necheance reelle",
        "Assurance\ntype", "Assurance\ndébut", "Assurance\nfin", "Statut", "PPSPS", "Contact",
        "Code de securite\nURSSAF", "Verification URSSAF\ndu 25/08/2026",
    ]
    for i, h in enumerate(entetes, start=1):
        ws.cell(row=4, column=i, value=h)
    lignes = [
        ("DP", "01 DESAMIANTAGE", "REVALECO", datetime(2026, 10, 5), datetime(2027, 2, 28), "OK", "MOCKVALIDE00001",
         "Attestation du 05/10/2026 - SIRET 12345678900012 - code à vérifier sur urssaf.fr"),
        ("PC", "01 DESAMIANTAGE", "REVALECO", datetime(2026, 10, 5), datetime(2027, 2, 28), "OK", "MOCKVALIDE00001",
         "Attestation du 05/10/2026 - SIRET 12345678900012 - code à vérifier sur urssaf.fr"),
        ("PC", "02 RAVALEMENT", "REUX", datetime(2026, 5, 12), datetime(2026, 9, 30), "PÉRIMÉE", "MOCKPERIMEE0001",
         "Authentique et valide - SIRET 90645051500059"),
        ("PC", "03 CHARPENTE", "INCONNU SARL", datetime(2026, 7, 1), datetime(2026, 12, 31), "OK", "MOCKINCONNU0001", None),
        ("PC", "13 ASCENSEUR", "(non attribué)", None, None, "URSSAF manquante", None, None),
    ]
    for r, (op, lot, ent, d_att, ech, statut, code, verif) in enumerate(lignes, start=5):
        ws.cell(row=r, column=1, value=op)
        ws.cell(row=r, column=2, value=lot)
        ws.cell(row=r, column=3, value=ent)
        ws.cell(row=r, column=4, value=d_att)
        ws.cell(row=r, column=5, value=ech)
        ws.cell(row=r, column=9, value=statut)
        ws.cell(row=r, column=12, value=code)
        ws.cell(row=r, column=13, value=verif)
    ws.cell(row=11, column=1, value="Nota : ligne de commentaire à ignorer")
    p = tmp_path / "suivi.xlsx"
    wb.save(p)
    return p
