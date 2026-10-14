"""Lecture et mise à jour du tableau Excel « Suivi des attestations par entreprise ».

Structure attendue (feuille « Suivi attestations ») :
- lignes de titre/sous-titre, puis une ligne d'en-têtes (détectée automatiquement) ;
- une ligne par (opération, lot, entreprise) ; les notas en bas de feuille sont ignorés
  (colonne Lot ou Entreprise vide).
Les colonnes sont repérées par leur en-tête (recherche insensible à la casse et aux retours
à la ligne), pas par leur lettre, pour résister aux ajouts de colonnes.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

FEUILLE_DEFAUT = "Suivi attestations"

# en-tête logique -> motifs acceptés (regex, insensible à la casse, sur l'en-tête sans retours à la ligne)
COLONNES = {
    "operation": r"^op[ée]ration",
    "lot": r"^lot",
    "entreprise": r"^entreprise",
    "date_attestation": r"urssaf.*date",
    "echeance": r"urssaf.*[ée]ch[ée]ance",
    "statut": r"^statut",
    "code": r"code de s[ée]curit[ée]",
    "verification": r"^v[ée]rification urssaf",
}


@dataclass
class LigneSuivi:
    ligne: int
    operation: str | None
    lot: str | None
    entreprise: str | None
    date_attestation: date | None
    echeance: date | None
    statut: str | None
    code_securite: str | None
    verification: str | None

    @property
    def code_a_verifier(self) -> bool:
        """Vrai tant que la colonne Vérification ne mentionne pas un contrôle effectué."""
        v = (self.verification or "").lower()
        if not self.code_securite:
            return False
        if re.search(r"[àa] v[ée]rifier", v):
            return True
        return not re.search(r"authentique|v[ée]rifi[ée]e? (sur|par)|p[ée]rim[ée]e selon|code inconnu", v)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["date_attestation"] = self.date_attestation.isoformat() if self.date_attestation else None
        d["echeance"] = self.echeance.isoformat() if self.echeance else None
        d["code_a_verifier"] = self.code_a_verifier
        return d


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and (m := re.match(r"(\d{2})/(\d{2})/(\d{4})", v.strip())):
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


class Suivi:
    def __init__(self, fichier: str | Path, feuille: str = FEUILLE_DEFAUT):
        self.fichier = Path(fichier)
        self.wb = load_workbook(self.fichier)
        if feuille not in self.wb.sheetnames:
            raise ValueError(f"Feuille {feuille!r} introuvable dans {self.fichier.name} ({self.wb.sheetnames})")
        self.ws: Worksheet = self.wb[feuille]
        self.ligne_entetes = 0
        self.colonnes = self._reperer_colonnes()

    def _reperer_colonnes(self) -> dict[str, int]:
        """Cherche la ligne d'en-têtes dans les 10 premières lignes (celle qui contient « Entreprise »)."""
        for ligne in range(1, min(10, self.ws.max_row) + 1):
            entetes = {c.column: _norm(c.value) for c in self.ws[ligne] if c.value}
            trouve: dict[str, int] = {}
            for nom, motif in COLONNES.items():
                for col, txt in entetes.items():
                    if re.search(motif, txt):
                        trouve[nom] = col
                        break
            if all(n in trouve for n in ("entreprise", "code", "verification")):
                self.ligne_entetes = ligne
                return trouve
        raise ValueError(
            f"Ligne d'en-têtes introuvable dans {self.fichier.name} : il faut au moins les colonnes "
            "« Entreprise », « Code de sécurité URSSAF » et « Vérification URSSAF »"
        )

    def _val(self, ligne: int, nom: str):
        col = self.colonnes.get(nom)
        return self.ws.cell(row=ligne, column=col).value if col else None

    def lignes(self) -> list[LigneSuivi]:
        out = []
        for r in range(self.ligne_entetes + 1, self.ws.max_row + 1):
            entreprise = self._val(r, "entreprise")
            lot = self._val(r, "lot")
            if not entreprise or not lot:
                continue  # notas, lignes vides
            code = self._val(r, "code")
            out.append(
                LigneSuivi(
                    ligne=r,
                    operation=_s(self._val(r, "operation")),
                    lot=_s(lot),
                    entreprise=_s(entreprise),
                    date_attestation=_date(self._val(r, "date_attestation")),
                    echeance=_date(self._val(r, "echeance")),
                    statut=_s(self._val(r, "statut")),
                    code_securite=re.sub(r"\s+", "", str(code)).upper() if code else None,
                    verification=_s(self._val(r, "verification")),
                )
            )
        return out

    def codes_uniques(self, seulement_a_verifier: bool = False) -> dict[str, list[LigneSuivi]]:
        """Code de sécurité -> lignes qui le portent (une entreprise sur plusieurs lots partage son code)."""
        par_code: dict[str, list[LigneSuivi]] = {}
        for l in self.lignes():
            if not l.code_securite:
                continue
            if seulement_a_verifier and not l.code_a_verifier:
                continue
            par_code.setdefault(l.code_securite, []).append(l)
        return par_code

    def ecrire_verification(self, ligne: int, texte: str, echeance: date | None = None, remplacer: bool = False) -> None:
        """Écrit le résultat de vérification (colonne Vérification) ; conserve l'historique sauf `remplacer`."""
        cell = self.ws.cell(row=ligne, column=self.colonnes["verification"])
        ancien = _s(cell.value) or ""
        if remplacer or not ancien:
            cell.value = texte
        else:
            # remplace la mention « code à vérifier … » par le résultat, sinon ajoute
            nettoye = re.sub(r"\s*-\s*code [àa] v[ée]rifier sur urssaf\.fr\.?", "", ancien, flags=re.I).strip()
            cell.value = f"{nettoye} - {texte}" if nettoye else texte
        if echeance and "echeance" in self.colonnes:
            self.ws.cell(row=ligne, column=self.colonnes["echeance"]).value = datetime(echeance.year, echeance.month, echeance.day)

    def enregistrer(self, destination: str | Path | None = None) -> Path:
        dest = Path(destination) if destination else self.fichier
        self.wb.save(dest)
        return dest


def _s(v) -> str | None:
    return None if v is None else str(v).strip()


def texte_resultat(res, controle=None, aujourdhui: date | None = None) -> str:
    """Phrase standard pour la colonne Vérification, dans le style déjà utilisé dans le tableau."""
    aujourdhui = aujourdhui or date.today()
    j = aujourdhui.strftime("%d/%m/%Y")
    if res.statut == "valide":
        sirets = res.sirets
        s = f"Authentique et valide - SIRET {sirets[0]}" if sirets else "Authentique et valide"
        if len(sirets) > 1:
            s += f" et {len(sirets) - 1} autre(s) etablissement(s)"
        s += f" (verifie par API Urssaf le {j})"
        if controle and controle.anomalies:
            s += " ; ANOMALIES : " + " ; ".join(controle.anomalies)
        return s
    if res.statut == "perimee":
        return f"PERIMEE selon l'Urssaf depuis le {res.perimee_depuis or '?'} (verifie par API le {j}) - attestation a jour a reclamer"
    if res.statut == "inconnue":
        return f"CODE INCONNU de l'Urssaf (verifie par API le {j}) - document suspect, a reclamer directement a l'entreprise"
    return f"Verification API Urssaf le {j} : {res.message}"
